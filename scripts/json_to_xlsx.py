#!/usr/bin/env python3
"""boats.json（source of truth）から data/master_kanagawa.xlsx を生成する。"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOATS_JSON = os.path.join(ROOT, "data", "boats.json")
XLSX_OUT = os.path.join(ROOT, "data", "master_kanagawa.xlsx")

HEADERS = [
    "No.", "都道府県", "市町村", "港", "船宿名", "URL", "業態",
    "釣果（直近）", "料金", "状態", "情報源", "収集日", "失敗回数",
]

STATUS_FILL = {
    "済": "C6EFCE",
    "一部": "FFEB9C",
    "取得不可": "FFC7CE",
    "未収集": "F2F2F2",
}


def fmt_catches(catches):
    lines = []
    for c in catches or []:
        parts = [c.get("fish", "")]
        if c.get("size"):
            parts.append(c["size"])
        if c.get("count"):
            parts.append(c["count"])
        if c.get("date"):
            parts.append(f"({c['date']})")
        line = " ".join(p for p in parts if p)
        if c.get("note"):
            line += f" ※{c['note']}"
        lines.append(line)
    return "\n".join(lines)


def fmt_prices(prices):
    lines = []
    for p in prices or []:
        yen = f"{p['yen']:,}円" if p.get("yen") is not None else "－"
        line = f"{p.get('label', '')} {yen}"
        if p.get("note"):
            line += f" ※{p['note']}"
        lines.append(line.strip())
    return "\n".join(lines)


def main():
    with open(BOATS_JSON, encoding="utf-8") as f:
        boats = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "神奈川県マスター"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for r, b in enumerate(boats, 2):
        row = [
            b.get("no"), b.get("pref", ""), b.get("city", ""), b.get("port", ""),
            b.get("name", ""), b.get("url", ""), b.get("type", ""),
            fmt_catches(b.get("catches")), fmt_prices(b.get("prices")),
            b.get("status", ""), b.get("source", ""), b.get("collected_at", ""),
            b.get("error_count", 0),
        ]
        for col, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (8, 9))
        fill = STATUS_FILL.get(b.get("status", ""))
        if fill:
            ws.cell(row=r, column=10).fill = PatternFill("solid", fgColor=fill)

    widths = [5, 9, 10, 10, 16, 34, 9, 46, 34, 8, 12, 11, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(XLSX_OUT)
    print(f"wrote {XLSX_OUT} ({len(boats)} boats)")


if __name__ == "__main__":
    main()
